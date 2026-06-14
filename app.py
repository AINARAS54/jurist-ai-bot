import os
import re
import logging
import base64
from io import BytesIO
from datetime import datetime, timezone, timedelta

import requests
from flask import Flask, request, jsonify
from supabase import create_client, Client
from pypdf import PdfReader
from docx import Document

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("justice_ai_pro")

app = Flask(__name__)

TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY", "").strip()
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o-mini").strip()
SUPABASE_URL = os.getenv("SUPABASE_URL", "").strip()
SUPABASE_SERVICE_ROLE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY", "").strip()
PUBLIC_BASE_URL = os.getenv("PUBLIC_BASE_URL", "").strip() or os.getenv("RENDER_EXTERNAL_URL", "").strip()
TEST_MODE = False  # production: Telegram Stars payments enabled

TELEGRAM_API = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}" if TELEGRAM_BOT_TOKEN else ""

if not TELEGRAM_BOT_TOKEN:
    raise RuntimeError("Missing TELEGRAM_BOT_TOKEN")
if not OPENAI_API_KEY:
    raise RuntimeError("Missing OPENAI_API_KEY")
if not SUPABASE_URL:
    raise RuntimeError("Missing SUPABASE_URL")
if not SUPABASE_SERVICE_ROLE_KEY:
    raise RuntimeError("Missing SUPABASE_SERVICE_ROLE_KEY")

supabase: Client = create_client(SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY)

USER_STATES = {}
MEDIA_GROUPS = {}

PLANS = {
    "1m": {"months": 1, "stars": 100, "lt": "1 mėnuo", "en": "1 month", "no": "1 måned"},
    "3m": {"months": 3, "stars": 250, "lt": "3 mėnesiai", "en": "3 months", "no": "3 måneder"},
    "6m": {"months": 6, "stars": 450, "lt": "6 mėnesiai", "en": "6 months", "no": "6 måneder"},
    "12m": {"months": 12, "stars": 800, "lt": "12 mėnesių", "en": "12 months", "no": "12 måneder"},
}

COUNTRIES = {
    "lt": {"lt": "Lietuva", "en": "Lithuania", "no": "Litauen"},
    "no": {"lt": "Norvegija", "en": "Norway", "no": "Norge"},
    "uk": {"lt": "Jungtinė Karalystė", "en": "United Kingdom", "no": "Storbritannia"},
    "other": {"lt": "Kita šalis", "en": "Other country", "no": "Annet land"},
}

SYSTEM_PROMPT = """
You are Justice AI, a consumer rights and legal claims assistant.

Language:
- Reply only in the selected user language.
- Never mix languages.

Jurisdiction:
- Apply the selected country/jurisdiction.
- If selected jurisdiction is Norway, do not answer as if only Lithuanian law is available. Use Norwegian consumer law and relevant Norwegian criminal law sections where applicable.
- If selected jurisdiction is United Kingdom, use UK law. If selected jurisdiction is Lithuania, use Lithuanian law.
- Always include relevant legal acts with article/section numbers when possible.
- Never state that a crime definitely occurred.
- Use cautious wording such as: possibly applicable, may be relevant, may be assessed under.

Style:
- Use grammatically correct, natural and professional language.
- Avoid repetition.
- Use clear complete sentences.
- Do not use greetings such as Dear Sir/Madam, Gerbiamasis, Gerbiamoji, Gerb. pone, Gerb. ponia.
- Do not use the phrase [Paraiška pridedama].
- In Lithuanian legal text, avoid the word "normaliai". Use "įprastai", "pagal paskirtį" or "įprastomis naudojimo sąlygomis".
- Telegram answers must be concise. Avoid long reports unless the user explicitly requests a detailed report.
- Do not repeat the same URL. If a Google Drive or other link is included, show it only once.

Rules:
- Do not claim to be a lawyer.
- Do not suggest consulting lawyers or external specialists.
- Provide practical next steps yourself.
- Mention institutions only as concrete action targets.
- Separate facts from assumptions.
- Never copy recipient/authority contact details into claimant/sender signature.
"""

WELCOME_TEXT_LT = """⚖️ JUSTICE AI

Vartotojų teisių apsaugos ir teisinių pretenzijų sistema

Padedu spręsti:

📦 Internetinių pirkimų problemas
💳 Pinigų grąžinimo klausimus
📄 Pretenzijų ir skundų rengimą
⚖️ Sutarčių ir paslaugų ginčus
🚨 Galimus sukčiavimo atvejus"""

WELCOME_TEXT_EN = """⚖️ JUSTICE AI

Consumer Rights & Legal Claims Assistant

I can help with:

📦 Online purchase problems
💳 Refunds and payment disputes
📄 Complaints and claims
⚖️ Contract and service disputes
🚨 Possible fraud cases"""

WELCOME_TEXT_NO = """⚖️ JUSTICE AI

Assistent for forbrukerrettigheter og juridiske krav

Jeg kan hjelpe med:

📦 Problemer med nettkjøp
💳 Refusjon og betalingskonflikter
📄 Klager og krav
⚖️ Kontrakt- og tjenestetvister
🚨 Mulige svindelsaker"""

SAFETY_TEXT = {
    "lt": """🔐 Saugumas ir konfidencialumas

• Jūsų pateikta informacija naudojama tik bylos analizei ir dokumentų rengimui.

• Duomenys saugomi iki prenumeratos pabaigos.

• Pasibaigus prenumeratai duomenys automatiškai ištrinami.

✅ Suprantu ir sutinku""",

    "en": """🔐 Security and confidentiality

• Your information is used only for case analysis and document preparation.

• Data is stored until the subscription ends.

• After the subscription ends, the data is automatically deleted.

✅ I understand and agree""",

    "no": """🔐 Sikkerhet og konfidensialitet

• Informasjonen brukes kun til saksanalyse og dokumentforberedelse.

• Data lagres til abonnementet avsluttes.

• Etter at abonnementet avsluttes slettes dataene automatisk.

✅ Jeg forstår og godtar""",
}

TEXT = {
    "lt": {
        "accept": "✅ Suprantu ir sutinku",
        "choose_country": "🌍 Pasirinkite šalį:",
        "collect_case": "📋 Surinkime duomenis bylai\n\nTrumpai aprašykite situaciją arba įkelkite dokumentą.\n\nKuo daugiau informacijos pateiksite, tuo tikslesnė bus analizė.",
        "thinking": "⚖️ Analizuoju pateiktą informaciją...",
        "case_created": "📁 Byla sukurta\n\n🆔 Bylos Nr.: {case_number}",
        "initial": "🔎 Tai pirminis bylos vertinimas.\n\nNorint gauti išsamų vertinimą, veiksmų planą ir dokumentų projektus, pasirinkite prieigos laikotarpį.",
        "choose_plan": "🔓 Pasirinkite prieigos laikotarpį:",
        "active_until": "✅ Prieiga aktyvi iki:",
        "access_active": "✅ Prieiga aktyvuota.",
        "test_active": "🧪 TEST REŽIMAS: prieiga aktyvuota.",
        "after_full": "❓ Norite patikslinti bylą?\n\nGalite užduoti papildomą klausimą arba sugeneruoti dokumentą pagal šią bylą.",
        "ask_question": "❓ Užduoti klausimą",
        "generate_doc": "📄 Generuoti dokumentą",
        "new_case": "➕ Nauja byla",
        "my_cases": "📁 Mano bylos",
        "new_case_started": "➕ Pradėkime naują bylą. Ankstesnės bylos lieka išsaugotos ir jas galėsite tęsti vėliau.",
        "ask_prompt": "❓ Parašykite papildomą klausimą dėl šios bylos.",
        "choose_doc": "📄 Pasirinkite generuoti:",
        "no_case": "Pirmiausia aprašykite situaciją arba įkelkite dokumentą.",
        "doc_thinking": "⏳ Ruošiu dokumentą...",
        "file_received": "📎 Failas gautas. Apdoroju dokumentą...",
        "file_saved": "📎 Dokumentas pridėtas prie bylos.",
        "file_no_text": "⚠️ Failą išsaugojau, bet teksto nuskaityti nepavyko. Trumpai aprašykite, kas jame svarbu.",
    },
    "en": {
        "accept": "✅ I understand and agree",
        "choose_country": "🌍 Select country:",
        "collect_case": "📋 Let us collect case details\n\nBriefly describe the situation or upload a document.\n\nThe more information you provide, the more accurate the assessment will be.",
        "thinking": "⚖️ Analyzing the submitted information...",
        "case_created": "📁 Case created\n\n🆔 Case No.: {case_number}",
        "initial": "🔎 This is an initial case assessment.\n\nTo receive a full review, action plan and document drafts, choose an access period.",
        "choose_plan": "🔓 Choose access period:",
        "active_until": "✅ Access active until:",
        "access_active": "✅ Access activated.",
        "test_active": "🧪 TEST MODE: access activated.",
        "after_full": "❓ Would you like to clarify the case?\n\nYou can ask a follow-up question or generate a document based on this case.",
        "ask_question": "❓ Ask a question",
        "generate_doc": "📄 Generate document",
        "new_case": "➕ New case",
        "my_cases": "📁 My cases",
        "new_case_started": "➕ Let us start a new case. Previous cases remain saved and can be continued later.",
        "ask_prompt": "❓ Write your follow-up question about this case.",
        "choose_doc": "📄 Select to generate:",
        "no_case": "Please describe the situation or upload a document first.",
        "doc_thinking": "⏳ Preparing document...",
        "file_received": "📎 File received. Processing the document...",
        "file_saved": "📎 Document added to the case.",
        "file_no_text": "⚠️ I saved the file, but could not extract text. Briefly describe what is important in it.",
    },
    "no": {
        "accept": "✅ Jeg forstår og godtar",
        "choose_country": "🌍 Velg land:",
        "collect_case": "📋 La oss samle saksinformasjon\n\nBeskriv situasjonen kort eller last opp et dokument.\n\nJo mer informasjon du gir, desto mer presis blir vurderingen.",
        "thinking": "⚖️ Analyserer innsendt informasjon...",
        "case_created": "📁 Sak opprettet\n\n🆔 Saksnr.: {case_number}",
        "initial": "🔎 Dette er en første vurdering av saken.\n\nFor full vurdering, handlingsplan og dokumentutkast, velg tilgangsperiode.",
        "choose_plan": "🔓 Velg tilgangsperiode:",
        "active_until": "✅ Tilgang aktiv til:",
        "access_active": "✅ Tilgang aktivert.",
        "test_active": "🧪 TESTMODUS: tilgang aktivert.",
        "after_full": "❓ Vil du avklare saken videre?\n\nDu kan stille et oppfølgingsspørsmål eller generere et dokument basert på saken.",
        "ask_question": "❓ Still spørsmål",
        "generate_doc": "📄 Generer dokument",
        "new_case": "➕ Ny sak",
        "my_cases": "📁 Mine saker",
        "new_case_started": "➕ La oss starte en ny sak. Tidligere saker lagres og kan fortsettes senere.",
        "ask_prompt": "❓ Skriv oppfølgingsspørsmålet ditt om denne saken.",
        "choose_doc": "📄 Velg å generere:",
        "no_case": "Beskriv situasjonen eller last opp et dokument først.",
        "doc_thinking": "⏳ Forbereder dokument...",
        "file_received": "📎 Fil mottatt. Behandler dokumentet...",
        "file_saved": "📎 Dokumentet er lagt til saken.",
        "file_no_text": "⚠️ Filen er lagret, men tekst kunne ikke leses. Beskriv kort hva som er viktig i den.",
    },
}


def detect_lang(user: dict) -> str:
    code = (user.get("language_code") or "").lower()
    if code.startswith("lt"):
        return "lt"
    if code.startswith("no") or code.startswith("nb") or code.startswith("nn"):
        return "no"
    return "en"


def language_for_country_selection(country: str, current_lang: str | None, tg_user: dict | None = None) -> str:
    """
    Country selection changes only the legal jurisdiction.
    The bot interface language stays the same as the current/Telegram language.
    """
    if current_lang in ("lt", "no", "en"):
        return current_lang

    tg_lang = detect_lang(tg_user or {})
    return tg_lang if tg_lang in ("lt", "no", "en") else "en"


def welcome_text(lang: str) -> str:
    if lang == "lt":
        return WELCOME_TEXT_LT
    if lang == "no":
        return WELCOME_TEXT_NO
    return WELCOME_TEXT_EN


def lang_name(lang: str) -> str:
    return {"lt": "Lithuanian", "no": "Norwegian", "en": "English"}.get(lang, "English")


def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def iso(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).isoformat()


def parse_dt(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except Exception:
        return None


def add_months(months: int) -> datetime:
    return now_utc() + timedelta(days=30 * months)


def get_state(chat_id: int) -> dict:
    if chat_id not in USER_STATES:
        USER_STATES[chat_id] = {
            "lang": "en",
            "accepted": False,
            "country": None,
            "case_text": None,
            "case_id": None,
            "case_number": None,
            "awaiting_followup": False,
            "awaiting_doc_details": False,
            "pending_doc_type": None,
            "doc_extra_data": "",
            "files": [],
            "upload_menu_sent": False,
        }
    return USER_STATES[chat_id]


def send_message(chat_id: int, text: str, reply_markup: dict | None = None):
    payload = {"chat_id": chat_id, "text": text, "disable_web_page_preview": True}
    if reply_markup:
        payload["reply_markup"] = reply_markup
    r = requests.post(f"{TELEGRAM_API}/sendMessage", json=payload, timeout=30)
    if not r.ok:
        logger.error("Telegram send error: %s %s", r.status_code, r.text)
    return r.json() if r.ok else None


def send_chunks(chat_id: int, text: str, reply_markup: dict | None = None):
    if not text:
        return
    chunks = [text[i:i + 3900] for i in range(0, len(text), 3900)]
    for idx, chunk in enumerate(chunks):
        send_message(chat_id, chunk, reply_markup if idx == len(chunks) - 1 else None)


def answer_callback(callback_query_id: str, text: str = ""):
    requests.post(f"{TELEGRAM_API}/answerCallbackQuery", json={"callback_query_id": callback_query_id, "text": text}, timeout=10)


def safety_menu(lang: str):
    return {"inline_keyboard": [[{"text": TEXT[lang]["accept"], "callback_data": "accept_safety"}]]}


def country_menu():
    return {"inline_keyboard": [
        [{"text": "🇱🇹 Lietuva", "callback_data": "country_lt"}],
        [{"text": "🇳🇴 Norvegija", "callback_data": "country_no"}],
        [{"text": "🇬🇧 Jungtinė Karalystė", "callback_data": "country_uk"}],
        [{"text": "🌐 Kita šalis", "callback_data": "country_other"}],
    ]}


def plan_menu(lang: str):
    labels = []
    for key, plan in PLANS.items():
        labels.append([{"text": f"⭐ {plan['stars']} - {plan[lang]}", "callback_data": f"plan_{key}"}])
    return {"inline_keyboard": labels}


def send_subscription_required(chat_id: int, lang: str):
    if lang == "lt":
        text = (
            "🔒 Prieš siunčiant bylos informaciją ar dokumentus reikia aktyvuoti prieigą.\n\n"
            "Pasirinkite abonementą:"
        )
    elif lang == "no":
        text = (
            "🔒 Før du sender saksinformasjon eller dokumenter, må du aktivere tilgang.\n\n"
            "Velg abonnement:"
        )
    else:
        text = (
            "🔒 Before sending case information or documents, access must be activated.\n\n"
            "Choose a subscription:"
        )
    send_message(chat_id, text, reply_markup=plan_menu(lang))


def after_full_menu(lang: str):
    return {"inline_keyboard": [
        [{"text": TEXT[lang]["ask_question"], "callback_data": "ask_question"}],
        [{"text": TEXT[lang]["generate_doc"], "callback_data": "show_docs"}],
        [{"text": TEXT[lang]["my_cases"], "callback_data": "case_list"}],
        [{"text": TEXT[lang]["new_case"], "callback_data": "new_case"}],
    ]}


def file_action_menu(lang: str):
    if lang == "lt":
        return {"inline_keyboard": [
            [{"text": "📂 Bylos dokumentai", "callback_data": "case_docs"}],
            [{"text": "✅ Analizuoti bylą", "callback_data": "analyze_case"}],
            [{"text": TEXT[lang]["my_cases"], "callback_data": "case_list"}],
            [{"text": TEXT[lang]["new_case"], "callback_data": "new_case"}],
        ]}
    if lang == "no":
        return {"inline_keyboard": [
            [{"text": "📂 Saksdokumenter", "callback_data": "case_docs"}],
            [{"text": "✅ Analyser saken", "callback_data": "analyze_case"}],
            [{"text": TEXT[lang]["my_cases"], "callback_data": "case_list"}],
            [{"text": TEXT[lang]["new_case"], "callback_data": "new_case"}],
        ]}
    return {"inline_keyboard": [
        [{"text": "📂 Case documents", "callback_data": "case_docs"}],
        [{"text": "✅ Analyze case", "callback_data": "analyze_case"}],
        [{"text": TEXT[lang]["my_cases"], "callback_data": "case_list"}],
        [{"text": TEXT[lang]["new_case"], "callback_data": "new_case"}],
    ]}

def detect_case_type(case_text: str) -> str:
    text = (case_text or "").lower()
    fraud_keywords = [
        "sukči", "apgav", "scam", "fraud", "deep", "coin", "crypto", "kript", "binance",
        "invest", "whatsapp", "telegram", "anydesk", "platform", "netikra", "praradau pinigus"
    ]
    bank_keywords = ["bank", "kortel", "visa", "mastercard", "mokėj", "paved", "chargeback", "lėšų grąž"]
    seller_keywords = ["prek", "pardav", "užsak", "siunt", "nepristat", "negav", "brok", "garant", "defekt", "grąžinti"]
    service_keywords = ["sutart", "paslaug", "prenumer", "operator", "internetas", "ryšys"]

    if any(k in text for k in fraud_keywords):
        return "fraud"
    if any(k in text for k in bank_keywords):
        return "bank"
    if any(k in text for k in seller_keywords):
        return "seller"
    if any(k in text for k in service_keywords):
        return "service"
    return "general"


def get_legal_reference_block(country: str, case_type: str, case_text: str, lang: str) -> str:
    """
    Fixed legal map. This prevents random legal sections from changing between analyses.
    """
    text = (case_text or "").lower()
    is_police = any(k in text for k in [
        "policija", "politiet", "politi", "henlegg", "nutrauk", "klage", "skund",
        "saksbehandlingskapasitet", "etterforskning", "tyrimą", "tyrimas"
    ])
    has_large_loss = any(k in text for k in [
        "18766", "18,766", "didel", "betydelig", "large loss", "økonomisk tap",
        "finansinis nuostolis", "financial loss"
    ])
    has_organized = any(k in text for k in [
        "organizuot", "organisert", "systemat", "whatsapp", "telegram", "platform", "platforma"
    ])

    if country == "no":
        if lang == "lt":
            lines = [
                "Naudok tik šiuos teisės aktus. Negeneruok kitų straipsnių.",
                "",
                "1. Straffeloven §371 – Bedrageri / sukčiavimas",
                "Konkreti dalis: §371 a punktas.",
                "Atitikimo logika: galimas suklaidinimas, dėl kurio asmuo atliko veiksmą ir patyrė nuostolį arba nuostolio riziką.",
                "Bylos faktai, kurie gali pagrįsti: finansinis nuostolis, kripto pervedimai, investicinė platforma, wallet / blockchain duomenys, susirašinėjimas su įtariamaisiais.",
            ]
            if has_large_loss or has_organized:
                lines += [
                    "",
                    "2. Straffeloven §372 – Grovt bedrageri / stambus sukčiavimas",
                    "Konkrečios dalys: §372 a punktas – reikšminga ekonominė žala; §372 c punktas – veika per kelis atvejus arba ilgesnį laiką; §372 d punktas – kelių asmenų bendras, sisteminis arba organizuotas pobūdis.",
                    "Naudok tik kaip 'gali būti aktualu', jeigu dokumentuose yra didelis nuostolis, tęstinė schema arba organizuotumo požymiai.",
                ]
            if is_police:
                lines += [
                    "",
                    "3. Straffeprosessloven – skundo dėl tyrimo nutraukimo / policijos sprendimo procedūra",
                    "Konkreti situacijos dalis: policijos sprendimas nutraukti tyrimą, skundas dėl nutraukimo, prašymas peržiūrėti sprendimą arba pateikti papildomus įrodymus.",
                    "Nerašyk konkretaus Straffeprosessloven paragrafo, jei jo tiksliai nėra byloje ar šiame fiksuotame sąraše.",
                ]
            return "\n".join(lines)

        if lang == "no":
            lines = [
                "Bruk bare disse rettsgrunnlagene. Ikke lag andre paragrafnumre.",
                "",
                "1. Straffeloven §371 – Bedrageri",
                "Relevant del: §371 bokstav a.",
                "Trefflogikk: mulig villfarelse som har fått personen til å handle og som har ført til tap eller fare for tap.",
                "Saksfakta som kan støtte dette: økonomisk tap, kryptotransaksjoner, investeringsplattform, wallet-/blockchain-dokumentasjon, kommunikasjon med mistenkte.",
            ]
            if has_large_loss or has_organized:
                lines += [
                    "",
                    "2. Straffeloven §372 – Grovt bedrageri",
                    "Relevante deler: §372 bokstav a – betydelig økonomisk skade; bokstav c – flere anledninger eller over lengre tid; bokstav d – flere i fellesskap eller systematisk/organisert preg.",
                    "Bruk bare forsiktig: 'kan være relevant' hvis dokumentene viser stort tap, langvarig eller organisert fremgangsmåte.",
                ]
            if is_police:
                lines += [
                    "",
                    "3. Straffeprosessloven – klage over henleggelse / politiets avgjørelse",
                    "Relevant situasjonsdel: politiets henleggelse, klage på henleggelsen, krav om ny vurdering eller innsending av nye bevis.",
                    "Ikke oppgi konkret paragrafnummer hvis det ikke fremgår sikkert av saken eller denne faste listen.",
                ]
            return "\n".join(lines)

        lines = [
            "Use only these legal references. Do not invent other sections.",
            "",
            "1. Straffeloven §371 – Fraud",
            "Relevant part: §371(a).",
            "Matching logic: possible deception causing the person to act, causing loss or risk of loss.",
            "Supporting facts: financial loss, cryptocurrency transfers, investment platform, wallet/blockchain evidence, communications with suspects.",
        ]
        if has_large_loss or has_organized:
            lines += [
                "",
                "2. Straffeloven §372 – Aggravated fraud",
                "Relevant parts: §372(a) significant financial damage; §372(c) repeated or longer period; §372(d) joint/systematic/organised character.",
                "Use only cautiously as 'may be relevant' if documents show major loss or organised/repeated conduct.",
            ]
        if is_police:
            lines += [
                "",
                "3. Straffeprosessloven – complaint against case closure / police decision",
                "Matching situation: police case closure, complaint/appeal, request for renewed assessment or submission of new evidence.",
                "Do not provide a specific section number unless certain from the case or this fixed list.",
            ]
        return "\n".join(lines)

    if lang == "lt":
        return (
            "Naudok tik teisės aktus, kurie aiškiai susiję su pasirinkta jurisdikcija ir bylos faktais. "
            "Kiekvienam teisės aktui nurodyk straipsnį / dalį, atitikimo lygį ir konkrečią bylos dalį. "
            "Jei konkretaus straipsnio nežinai tiksliai, nerašyk numerio."
        )
    if lang == "no":
        return (
            "Bruk bare rettsgrunnlag som tydelig passer jurisdiksjonen og saksfakta. "
            "For hvert rettsgrunnlag skal du vise paragraf/del, treffnivå og konkret del av saken. "
            "Hvis paragrafnummeret ikke er sikkert, ikke oppgi nummer."
        )
    return (
        "Use only legal references clearly matching the jurisdiction and case facts. "
        "For each reference, show section/part, match level and the concrete part of the case. "
        "If the exact section number is not certain, do not provide a number."
    )


def normalize_case_answer(text: str) -> str:
    if not text:
        return text
    text = text.strip()
    headings = [
        "🆔 Bylos numeris", "📋 Situacija", "⚖️ Galimai taikytini teisės aktai",
        "🎯 Rekomenduojami veiksmai", "📄 Galimi dokumentai", "📄 Rekomenduojami dokumentai",
        "🆔 Case number", "📋 Situation", "⚖️ Potentially applicable legal references",
        "🎯 Recommended actions", "📄 Recommended documents",
        "🆔 Saksnummer", "📋 Situasjon", "⚖️ Mulig relevante rettsgrunnlag",
        "🎯 Anbefalte tiltak", "📄 Anbefalte dokumenter"
    ]
    for h in headings:
        text = re.sub(r"\s*" + re.escape(h), "\n\n" + h, text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def sanitize_generated_document(text: str) -> str:
    if not text:
        return text
    text = re.sub(r"^\s*(Sukūrė|Sukure|Created by|Opprettet av)\s+Justice AI\s*\n+", "", text, flags=re.IGNORECASE)
    text = re.sub(r"(?im)^\s*(ID|Asmens kodas|Personnummer|Fødselsnummer)\s*:\s*.*\n?", "", text)
    text = re.sub(r"\[[^\]]+\]", "", text)
    text = re.sub(r"(?im)^\s*(El\.?\s*paštas|E-?post|Email|Telefonas|Phone|Telefon|Adresas|Address)\s*:\s*$\n?", "", text)
    text = text.replace("Remiantis Straffeloven §371", "Galimai aktualus Straffeloven §371")
    text = text.replace("Remiantis Straffeloven", "Galimai aktualus Straffeloven")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def doc_label(doc_type: str, lang: str, country: str | None = None) -> str:
    country = country or "lt"
    labels = {
        "lt": {
            "seller_claim": "📄 Pretenzija pardavėjui",
            "bank_refund": "💳 Prašymas bankui dėl pinigų grąžinimo",
            "authority_complaint": "⚖️ Skundas institucijai",
            "fraud_report": "🚨 Pareiškimas dėl galimo sukčiavimo",
            "appeal_police_decision": "⚖️ Skundas dėl tyrimo nutraukimo / neveikimo",
        },
        "no": {
            "seller_claim": "📄 Klagebrev til selger",
            "bank_refund": "💳 Forespørsel til bank",
            "authority_complaint": "⚖️ Klage til myndighet",
            "fraud_report": "🚨 Svindelrapport",
            "appeal_police_decision": "⚖️ Klage på henleggelse / manglende svar",
        },
        "en": {
            "seller_claim": "📄 Seller complaint letter",
            "bank_refund": "💳 Bank refund request",
            "authority_complaint": "⚖️ Authority complaint",
            "fraud_report": "🚨 Fraud report",
            "appeal_police_decision": "⚖️ Appeal against case closure / no response",
        },
    }
    return labels.get(lang, labels["en"]).get(doc_type, doc_type)


def relevant_doc_types(case_text: str, country: str | None = None) -> list[str]:
    text = (case_text or "").lower()
    case_type = detect_case_type(case_text)

    police_context = any(k in text for k in [
        "policija", "politiet", "politi", "pareiškim", "pareiskim", "skund",
        "nutrauk", "henlegg", "atmet", "neatsak", "atsakymo negav", "bylos eiga"
    ])
    bank_context = any(k in text for k in [
        "bank", "kortel", "visa", "mastercard", "mokėj", "mokej", "paved",
        "chargeback", "lėšų", "lesu", "binance", "spectrocoin", "usdc", "crypto", "kript"
    ])

    if case_type == "fraud" or police_context:
        docs = []
        if police_context:
            docs.append("appeal_police_decision")
        else:
            docs.append("fraud_report")
        if bank_context:
            docs.append("bank_refund")
        return docs

    if case_type == "bank":
        return ["bank_refund", "authority_complaint"]
    if case_type == "seller":
        return ["seller_claim", "bank_refund", "authority_complaint"]
    if case_type == "service":
        return ["seller_claim", "authority_complaint"]
    return ["authority_complaint"]


def doc_menu(lang: str, case_text: str = "", country: str | None = None):
    rows = []
    for doc_type in relevant_doc_types(case_text, country):
        rows.append([{"text": doc_label(doc_type, lang, country), "callback_data": f"doc_{doc_type}"}])
    return {"inline_keyboard": rows}


def db_user_get(telegram_id: int):
    res = supabase.table("users").select("*").eq("telegram_id", telegram_id).execute()
    return res.data[0] if res.data else None


def db_user_upsert(tg_user: dict, lang: str | None = None):
    telegram_id = tg_user.get("id")
    username = tg_user.get("username")
    data = {"telegram_id": telegram_id, "username": username}
    if lang:
        data["language"] = lang
    existing = db_user_get(telegram_id)
    if existing:
        supabase.table("users").update(data).eq("telegram_id", telegram_id).execute()
    else:
        data["language"] = lang or "lt"
        supabase.table("users").insert(data).execute()


def db_set_subscription(telegram_id: int, plan_name: str, months: int, stars: int):
    until = add_months(months)
    supabase.table("users").update({"subscription_until": iso(until), "plan_name": plan_name}).eq("telegram_id", telegram_id).execute()
    supabase.table("payments").insert({"telegram_id": telegram_id, "plan_name": plan_name, "stars_paid": stars, "valid_until": iso(until)}).execute()
    return until


def db_subscription_active(telegram_id: int):
    user = db_user_get(telegram_id)
    if not user:
        return False, None
    until = parse_dt(user.get("subscription_until"))
    if not until:
        return False, None
    return until > now_utc(), until


def db_create_case(telegram_id: int, case_text: str):
    title = (case_text or "New case")[:80]
    res = supabase.table("cases").insert({"telegram_id": telegram_id, "title": title, "case_text": case_text, "status": "open"}).execute()
    case = res.data[0]
    case_id = case["id"]
    case_number = f"CASE-{now_utc().year}-{case_id:06d}"
    supabase.table("cases").update({"case_number": case_number}).eq("id", case_id).execute()
    return case_id, case_number


def db_update_case(case_id: int, case_text: str):
    if case_id:
        supabase.table("cases").update({"case_text": case_text}).eq("id", case_id).execute()


def db_add_case_file(case_id: int, file_name: str, file_type: str, telegram_file_id: str):
    if case_id:
        supabase.table("case_files").insert({"case_id": case_id, "file_name": file_name, "file_type": file_type, "telegram_file_id": telegram_file_id}).execute()


def db_list_cases(telegram_id: int, limit: int = 10):
    try:
        res = (
            supabase.table("cases")
            .select("id, case_number, title, status, created_at")
            .eq("telegram_id", telegram_id)
            .order("id", desc=True)
            .limit(limit)
            .execute()
        )
        return res.data or []
    except Exception as e:
        logger.warning("Case list failed: %s", e)
        return []


def db_get_case(case_id: int, telegram_id: int):
    try:
        res = (
            supabase.table("cases")
            .select("*")
            .eq("id", case_id)
            .eq("telegram_id", telegram_id)
            .execute()
        )
        return res.data[0] if res.data else None
    except Exception as e:
        logger.warning("Case get failed: %s", e)
        return None


def db_list_case_files(case_id: int):
    try:
        res = supabase.table("case_files").select("file_name").eq("case_id", case_id).execute()
        return [row.get("file_name") for row in (res.data or []) if row.get("file_name")]
    except Exception as e:
        logger.warning("Case file list failed: %s", e)
        return []


def extract_text_from_file(file_name: str, file_bytes: bytes) -> str:
    lower = (file_name or "").lower()
    try:
        if lower.endswith(".txt"):
            return file_bytes.decode("utf-8", errors="ignore")
        if lower.endswith(".pdf"):
            reader = PdfReader(BytesIO(file_bytes))
            text = "\\n".join((page.extract_text() or "") for page in reader.pages).strip()
            if len(text) >= 30:
                return text
            ocr_text = extract_pdf_with_openai_vision(file_bytes)
            return ocr_text.strip() if ocr_text else text
        if lower.endswith(".docx"):
            doc = Document(BytesIO(file_bytes))
            return "\n".join(p.text for p in doc.paragraphs).strip()
    except Exception as e:
        logger.warning("File extraction failed: %s", e)
    return ""


def extract_pdf_with_openai_vision(file_bytes: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except Exception as e:
        logger.warning("PyMuPDF not available for OCR fallback: %s", e)
        return ""

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        texts = []
        for page_index in range(min(len(doc), 2)):
            page = doc[page_index]
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            png_bytes = pix.tobytes("png")
            b64 = base64.b64encode(png_bytes).decode("utf-8")
            prompt = (
                "Extract all visible text from this document page. "
                "Return only the extracted text. Preserve names, dates, case numbers, addresses and email addresses."
            )
            r = requests.post(
                "https://api.openai.com/v1/chat/completions",
                headers={"Authorization": f"Bearer {OPENAI_API_KEY}", "Content-Type": "application/json"},
                json={
                    "model": OPENAI_MODEL,
                    "messages": [
                        {
                            "role": "user",
                            "content": [
                                {"type": "text", "text": prompt},
                                {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}},
                            ],
                        }
                    ],
                    "temperature": 0,
                },
                timeout=90,
            )
            if r.ok:
                texts.append(r.json()["choices"][0]["message"]["content"].strip())
            else:
                logger.warning("OpenAI Vision OCR failed: %s %s", r.status_code, r.text[:500])
        return "\n\n".join(t for t in texts if t)
    except Exception as e:
        logger.warning("PDF OCR fallback failed: %s", e)
        return ""


def openai_request(prompt: str, lang: str) -> str:
    r = requests.post(
        "https://api.openai.com/v1/chat/completions",
        headers={"Authorization": f"Bearer {OPENAI_API_KEY}", "Content-Type": "application/json"},
        json={
            "model": OPENAI_MODEL,
            "messages": [
                {"role": "system", "content": SYSTEM_PROMPT + f"\nReply only in {lang_name(lang)}."},
                {"role": "user", "content": prompt},
            ],
            "temperature": 0.2,
        },
        timeout=60,
    )
    if not r.ok:
        logger.error("OpenAI error: %s %s", r.status_code, r.text)
        return ""
    return r.json()["choices"][0]["message"]["content"].strip()


def build_case_prompt(case_text: str, lang: str, country: str, full: bool):
    country_name = COUNTRIES.get(country, COUNTRIES["lt"])[lang]
    case_type = detect_case_type(case_text)
    legal_block = get_legal_reference_block(country, case_type, case_text, lang)

    if lang == "lt":
        if full:
            return f"""
Atsakyk tik lietuviškai. Taikoma šalis / jurisdikcija: {country_name}.

Paruošk pilną, bet aiškų bylos vertinimą Telegram formatui.

BŪTINAS FORMATAS:

🆔 Bylos numeris

...

📋 Situacija

...

⚖️ Galimai taikytini teisės aktai

1. [teisės aktas]

📌 Konkreti dalis
...

📊 Atitikimas
Tiesioginis / stiprus / galimas atitikimas (procentas)

📄 Atitinkanti bylos dalis
...

🎯 Rekomenduojami veiksmai

1. ...
2. ...
3. ...

📄 Rekomenduojami dokumentai

• ...
• ...

Taisyklės:
- Eilučių skaičius neribojamas.
- Visada palik tuščią eilutę tarp skyrių.
- Nenaudok vientiso sulipusio teksto bloko.
- Teisės aktų skyriuje naudok tik fiksuotą sąrašą žemiau.
- Negeneruok kitų straipsnių ar paragrafų, kurių nėra fiksuotame sąraše.
- Kiekvienam teisės aktui privalomai nurodyk konkrečią dalį/punktą, atitikimo lygį, procentą ir bylos faktą.
- Neteik kategoriško teiginio, kad nusikaltimas įvykdytas. Naudok: „galimai taikytina“, „gali būti aktualu“, „gali būti vertinama pagal“.
- „Galimi dokumentai“ vadink „Rekomenduojami dokumentai“ ir nurodyk dokumentus, kuriuos galima sugeneruoti, ne techninius failų pavadinimus.

Fiksuoti teisės aktai ir atitikimo logika:
{legal_block}

Bylos informacija:
{case_text}
"""
        return f"""
Atsakyk tik lietuviškai. Taikoma šalis / jurisdikcija: {country_name}.

Paruošk pirminį bylos vertinimą su tarpais tarp skyrių.

Naudok šiuos skyrius:
📋 Kategorija

⚠️ Pagrindinė problema

⚖️ Galimai taikytini teisės aktai

📂 Trūkstami įrodymai

🎯 Pirmas rekomenduojamas veiksmas

🔓 Išplėstinio atsakymo galimybė

Teisės aktus naudok tik iš šio fiksuoto sąrašo:
{legal_block}

Situacija:
{case_text}
"""

    if lang == "no":
        return f"""
Svar kun på norsk. Jurisdiksjon: {country_name}.

Lag en tydelig {'full vurdering' if full else 'førstevurdering'} for Telegram med blank linje mellom seksjoner.

Bruk bare disse rettsgrunnlagene. For hvert rettsgrunnlag skal du vise konkret del, treffnivå, prosent og hvilken del av saken som passer:
{legal_block}

Saksinformasjon:
{case_text}
"""

    return f"""
Answer only in English. Jurisdiction: {country_name}.

Prepare a clear {'full case review' if full else 'initial assessment'} for Telegram with blank lines between sections.

Use only these legal references. For each reference, show exact part, match level, percentage and the concrete part of the case that matches:
{legal_block}

Case information:
{case_text}
"""

def document_name(doc_type: str, lang: str) -> str:
    names = {
        "lt": {
            "seller_claim": "pretenziją pardavėjui",
            "bank_refund": "prašymą bankui dėl pinigų grąžinimo",
            "authority_complaint": "skundą institucijai",
            "fraud_report": "pareiškimą dėl galimo sukčiavimo",
            "appeal_police_decision": "skundą dėl tyrimo nutraukimo arba institucijos neveikimo",
        },
        "en": {
            "seller_claim": "seller complaint letter",
            "bank_refund": "bank refund request",
            "authority_complaint": "consumer authority complaint",
            "fraud_report": "fraud report",
            "appeal_police_decision": "appeal against case closure or authority inaction",
        },
        "no": {
            "seller_claim": "klagebrev til selger",
            "bank_refund": "forespørsel til bank om tilbakebetaling",
            "authority_complaint": "klage til myndighet",
            "fraud_report": "svindelrapport",
            "appeal_police_decision": "klage på henleggelse eller manglende svar fra myndighet",
        },
    }
    return names.get(lang, names["en"]).get(doc_type, "document")


def jurisdiction_recipient_hint(country: str, doc_type: str, lang: str) -> str:
    if country == "no":
        if doc_type == "appeal_police_decision":
            return "Norvegijos policija / Politiet arba aukštesnė policijos institucija pagal bylos dokumentus. Nenaudok Forbrukertilsynet, jei byla apie policijos sprendimą ar neatsakytą skundą."
        if doc_type == "fraud_report":
            return "Norvegijos policija / Politiet. Jei dokumentas skirtas vartotojų ginčui, naudok Forbrukertilsynet."
        if doc_type == "authority_complaint":
            return "Forbrukertilsynet (Norwegian Consumer Authority) arba Forbrukerrådet, pagal situaciją."
        if doc_type == "bank_refund":
            return "Kliento bankas Norvegijoje arba mokėjimo kortelės išdavėjas."
        return "Pardavėjas arba paslaugos teikėjas Norvegijoje."
    if country == "uk":
        if doc_type == "appeal_police_decision":
            return "UK Police / relevant police force complaints department, jei byla apie policijos sprendimą arba neatsakytą skundą."
        if doc_type == "fraud_report":
            return "Action Fraud / UK Police, jei yra sukčiavimo požymių."
        if doc_type == "authority_complaint":
            return "Trading Standards arba Citizens Advice Consumer Service, pagal situaciją."
        if doc_type == "bank_refund":
            return "Kliento bankas / kortelės išdavėjas Jungtinėje Karalystėje."
        return "Pardavėjas arba paslaugos teikėjas Jungtinėje Karalystėje."
    if country == "other":
        return "Parink instituciją pagal byloje nurodytą šalį. Jei šalies nėra, dokumente nenaudok konkrečios institucijos pavadinimo."
    # Lithuania default
    if doc_type == "appeal_police_decision":
        return "Policija / aukštesnė policijos institucija arba prokuratūra, jei byla apie nutrauktą tyrimą ar neatsakytą skundą."
    if doc_type == "fraud_report":
        return "Policija / ePolicija, jei yra galimo sukčiavimo požymių."
    if doc_type == "authority_complaint":
        return "Valstybinė vartotojų teisių apsaugos tarnyba, jei tai vartotojų ginčas; Lietuvos bankas, jei tai finansinių paslaugų ginčas."
    if doc_type == "bank_refund":
        return "Kliento bankas arba mokėjimo kortelės išdavėjas."
    return "Pardavėjas arba paslaugos teikėjas."


def build_doc_prompt(case_text: str, lang: str, country: str, doc_type: str):
    country_name = COUNTRIES.get(country, COUNTRIES["lt"])[lang]
    doc = document_name(doc_type, lang)
    recipient_hint = jurisdiction_recipient_hint(country, doc_type, lang)
    case_type = detect_case_type(case_text)
    legal_block = get_legal_reference_block(country, case_type, case_text, lang)

    if lang == "lt":
        return f"""
Atsakyk tik lietuviškai. Taikoma šalis / jurisdikcija: {country_name}.

Paruošk dokumentą: {doc}.
Bylos tipas pagal turinį: {case_type}.
Adresato parinkimo gairė: {recipient_hint}

Svarbiausios taisyklės:
- Dokumentą generuok iš realios bylos informacijos ir iš įkeltų failų nuskaityto teksto.
- Jei PDF / bylos tekste yra vardas, pavardė, adresas, el. paštas, telefonas, data, suma, platforma, įmonė, bankas, pavedimų duomenys ar kiti faktai, įrašyk juos į dokumentą automatiškai.
- Nenaudok tuščių laukų [Vardas Pavardė], [Adresas], [El. paštas], [Telefonas], [Data], jei duomenys jau randami byloje.
- Jei konkretaus duomens nėra, jo eilutę praleisk. Neįterpk instrukcinių tekstų laužtiniuose skliaustuose.
- Jeigu prieš dokumento generavimą vartotojas pateikė papildomus trūkstamus duomenis, naudok juos dokumente.
- Dokumento adresatą nustatyk automatiškai pagal dokumento tipą, pasirinktą šalį ir bylos turinį. Neklausk „Kam adresuoti dokumentą?“, nebent byloje visiškai neaišku, ar dokumentas turi būti skirtas pardavėjui, bankui, policijai ar institucijai.
- Jei byla susijusi su policijos atsakymu, nutrauktu tyrimu, skundu dėl bylos eigos ar neatsakytu el. laišku, adresatą rinkis policijos instituciją, o ne vartotojų teisių instituciją.
- Nenaudok frazių: [Čia pateikite...], [Jei yra priedų...], [Paraiška pridedama].
- Niekada nepalik tuščių laukų ar laužtinių skliaustų dokumente. Jei trūksta nebūtino duomens, jo lauką praleisk.
- Dokumento datą nustatyk automatiškai pagal šiandienos datą, jei konkreti dokumento data nenurodyta byloje.
- Justice AI vidinio bylos numerio dokumente nerašyk. Naudok tik oficialų policijos, banko ar institucijos bylos numerį, jeigu toks randamas dokumentuose.
- Dokumentas turi būti pilnai užpildytas ir paruoštas siuntimui.
- Dokumentą pradėk nuo tinkamo gavėjo pagal pasirinktą šalį ir dokumento antraštės.
- Jei pasirinkta Norvegija, nenaudok Lietuvos institucijų, nebent byloje aiškiai nurodyta, kad ginčas nagrinėjamas Lietuvoje.
- Jei pasirinkta Jungtinė Karalystė, nenaudok Lietuvos institucijų.
- Nenaudok kreipinių: Gerbiamasis, Gerbiamoji, Gerb. pone, Gerb. ponia.
- Dokumento viršuje nerašyk „Sukūrė Justice AI“. Dokumentą pradėk nuo gavėjo arba dokumento pavadinimo.
- Įtrauk tik fiksuotus galimai taikytinus teisės aktus iš žemiau esančio sąrašo. Prie kiekvieno nurodyk konkrečią dalį, atitikimo pagrindą ir bylos faktus. Neteik kategoriško teiginio, kad nusikaltimas įvykdytas.
- Jei dokumente yra interneto nuoroda, rodyk ją tik vieną kartą. Nekartok tos pačios nuorodos skliaustuose ir nekartok jos keliose vietose.
- Google Drive nuorodą pateik taip: „Google Drive: https://...“.
- Jei yra aiškūs priedai pagal bylą, naudok skyrių „Priedai:“ ir išvardink juos profesionaliai pagal turinį, ne techniniais failų pavadinimais. Pvz. „Policijos pranešimas apie tyrimo nutraukimą (2026-05-04)“.
- Jei nėra el. pašto, telefono ar adreso, nerodyk laukų [El. paštas], [Telefonas], [Adresas]. Trūkstamus laukus praleisk.
- Nesiūlyk kreiptis į konsultantus ar teisininkus.
- Griežtai atskirk adresato / institucijos kontaktus nuo pareiškėjo kontaktų.
- Niekada nenaudok institucijos telefono ar el. pašto kaip pareiškėjo telefono ar el. pašto.
- Jei el. paštas baigiasi @politiet.no, @police.uk, @gov.uk, @vvtat.lt ar priklauso institucijai, nelaikyk jo pareiškėjo el. paštu.
- Jei telefono numeris randamas šalia policijos, banko ar institucijos pavadinimo, nelaikyk jo pareiškėjo telefonu.
- Pareiškėjo paraše rodyk tik tuos kontaktus, kurie aiškiai priklauso pareiškėjui. Jei abejoji, telefono ir el. pašto nerodyk.
- Vardą ir pavardę formatuok natūraliai: „Ainaras Kalnenas“, ne „Ainaras KALNENAS“.
- Niekada automatiškai nepridėk ID / asmens kodo / fødselsnummer eilutės, nebent vartotojas aiškiai paprašė.

Fiksuoti teisės aktai ir atitikimo logika:
{legal_block}

Bylos informacija ir nuskaitytas dokumentų tekstas:
{case_text}
"""

    return f"""
Reply only in {lang_name(lang)}. Jurisdiction: {country_name}.

Prepare this document: {doc}.
Case type from content: {case_type}.
Recipient guidance: {recipient_hint}

Rules:
- Generate the document from the real case information and extracted file text.
- If the case/PDF contains name, address, email, phone, dates, amounts, platform names, company names, bank/payment data or other facts, insert them directly.
- Do not use empty placeholders if data exists.
- If a data point is missing, omit that line instead of adding instructional bracket text.
- If the user provided additional missing data before document generation, use it in the document.
- Determine the document recipient automatically from the document type, selected jurisdiction and case content. Do not ask who the document should be addressed to unless it is completely unclear whether it should go to a seller, bank, police or authority.
- If the case concerns police response, discontinued investigation, case status, complaint progress or unanswered email, choose the police authority as recipient, not a consumer authority.
- Do not use instructional phrases like [insert details here].
- Never leave placeholders or square brackets in the document. If non-essential data is missing, omit that line.
- Set the document date automatically to today if no specific document date is found in the case.
- Do not include the internal Justice AI case number in the document. Use only an official police, bank or authority case/reference number if it is found in the documents.
- The document must be complete and ready to send.
- Start with the correct recipient for the selected jurisdiction and the document title.
- If Norway is selected, do not use Lithuanian institutions unless the case explicitly concerns Lithuania.
- If the UK is selected, do not use Lithuanian institutions.
- Do not use Dear Sir/Madam or gendered salutations.
- Do not write "Created by Justice AI" at the top. Start with the recipient or document title.
- Include only fixed legal references from the list below. For each, show relevant part, matching reason and case facts. Do not state that a crime definitely occurred.
- If a URL appears in the document, show it only once. Do not repeat the same URL in brackets or in multiple places.
- Format Google Drive links as: "Google Drive: https://...".
- If real attachments are identifiable from the case, include an Attachments section and describe them professionally by content, not by raw technical filenames. Otherwise omit it.
- Do not show missing placeholders such as [Email], [Phone], [Address] if the data is not present in the case.
- Do not suggest external lawyers or advisors.
- Strictly separate recipient/authority contact details from claimant contact details.
- Never use an authority phone number or email address as the claimant's phone number or email address.
- If an email ends with @politiet.no, @police.uk, @gov.uk, @vvtat.lt or clearly belongs to an institution, do not treat it as the claimant's email.
- If a phone number appears near a police, bank or authority name, do not treat it as the claimant's phone number.
- In the claimant signature, include only contact details that clearly belong to the claimant. If unsure, omit phone and email.
- Format claimant names naturally, for example “Ainaras Kalnenas”, not “Ainaras KALNENAS”.
- Never automatically add ID / personal number / fødselsnummer unless the user explicitly asked for it.

Fixed legal references and matching logic:
{legal_block}

Case information and extracted document text:
{case_text}
"""


def build_doc_missing_prompt(case_text: str, lang: str, country: str, doc_type: str) -> str:
    country_name = COUNTRIES.get(country, COUNTRIES["lt"])[lang]
    doc = document_name(doc_type, lang)
    recipient_hint = jurisdiction_recipient_hint(country, doc_type, lang)

    if lang == "lt":
        return f"""
Atsakyk tik lietuviškai. Taikoma šalis / jurisdikcija: {country_name}.

Patikrink, ar pakanka duomenų pilnai užpildytam dokumentui: {doc}.
Adresatą nustatyk automatiškai pagal šias gaires: {recipient_hint}

Svarbu:
- Pirmiausia pats nustatyk adresatą iš bylos konteksto, dokumento tipo ir pasirinktos šalies.
- Neklausk „Kam adresuoti dokumentą?“, jei galima suprasti, ar dokumentas skirtas pardavėjui, bankui, policijai ar institucijai.
- Jei byla susijusi su policijos sprendimu, skundo eiga, neatsakytu el. laišku ar tyrimo nutraukimu, adresatas yra policijos institucija, ne vartotojų teisių institucija.
- Dokumento tikslą, prašymo esmę, datą, terminą ir adresatą pirmiausia nustatyk iš bylos ir PDF dokumentų. Neklausk apie juos, jei tai galima suprasti iš konteksto.
- Kontaktinis telefonas nėra būtinas, jei jo nėra byloje; jo neklausk, nebent dokumento tipui jis kritiškai būtinas.
- Jei trūksta tik nebūtinų duomenų, atsakyk READY.
- Klausimą užduok tik tada, kai trūksta pareiškėjo vardo / pavardės arba visiškai neįmanoma nustatyti oficialaus prašymo tikslo iš bylos.
- Klausimus užduok tik dėl tikrai svarbių duomenų, be kurių dokumentas būtų nepilnas arba netikslus.
- Maksimaliai 3 klausimai.
- Klausimai turi būti trumpi ir konkretūs.
- Neklausk abstrakčių klausimų kaip „kokia informacija jus domina“. Klausk konkrečių trūkstamų faktų.

Atsakyk tik vienu iš šių formatų:
READY

arba:
MISSING:
1. ...
2. ...
3. ...

Bylos informacija ir nuskaitytas dokumentų tekstas:
{case_text}
"""

    return f"""
Reply only in {lang_name(lang)}. Jurisdiction: {country_name}.

Check whether there is enough information to generate a complete document: {doc}.
Determine the recipient automatically using this guidance: {recipient_hint}

Important:
- First determine the recipient from case context, document type and selected jurisdiction.
- Do not ask who the document should be addressed to if it can be determined whether it goes to a seller, bank, police or authority.
- If the case concerns police decision, complaint progress, unanswered email or discontinued investigation, the recipient is a police authority, not a consumer authority.
- First infer document purpose, request, date/timeframe and recipient from the case and PDF documents. Do not ask about them if they can be understood from context.
- A phone number is not mandatory; do not ask for it unless it is critical for the document type.
- If only non-essential data is missing, answer READY.
- Ask only if the claimant full name is missing or if the official request purpose is impossible to determine from the case.
- Ask only for essential missing data without which the document would be incomplete or inaccurate.
- Maximum 3 questions.
- Questions must be short and specific.
- Do not ask abstract questions like “what information are you interested in”; ask for concrete missing facts.

Reply only in one of these formats:
READY

or:
MISSING:
1. ...
2. ...
3. ...

Case information and extracted document text:
{case_text}
"""


def parse_missing_doc_questions(answer: str) -> str:
    text = (answer or "").strip()
    if not text:
        return ""
    if text.upper().startswith("READY"):
        return ""
    if "MISSING:" in text.upper():
        parts = re.split(r"MISSING:\s*", text, flags=re.IGNORECASE, maxsplit=1)
        if len(parts) == 2:
            return parts[1].strip()
    return ""


def missing_doc_message(lang: str, questions: str) -> str:
    if lang == "lt":
        return "📋 Dokumentui parengti trūksta duomenų:\n\n" + questions + "\n\nAtsakykite viena žinute."
    if lang == "no":
        return "📋 Dokumentet mangler noen viktige opplysninger:\n\n" + questions + "\n\nSvar i én melding."
    return "📋 The document is missing a few important details:\n\n" + questions + "\n\nReply in one message."

def document_has_placeholders(text: str) -> bool:
    if not text:
        return False
    bad_patterns = [
        r"\[[^\]]+\]",
        r"Čia pateikite",
        r"Jei yra priedų",
        r"Vardas Pavardė",
        r"El\. paštas",
        r"Telefonas",
        r"Adresas",
        r"Data:\s*$",
        r"Priedai:\s*\n\s*Nėra",
    ]
    return any(re.search(pat, text, flags=re.IGNORECASE | re.MULTILINE) for pat in bad_patterns)


def placeholder_fix_message(lang: str) -> str:
    if lang == "lt":
        return (
            "📋 Dokumentui parengti trūksta svarbiausių duomenų:\n\n"
            "1. Vardas ir pavardė\n"
            "2. El. paštas arba kitas kontaktas, jei norite jį įtraukti\n\n"
            "Atsakykite viena žinute."
        )
    if lang == "no":
        return (
            "📋 Dokumentet mangler de viktigste opplysningene:\n\n"
            "1. Navn og etternavn\n"
            "2. E-post eller annen kontakt hvis du vil ta den med\n\n"
            "Svar i én melding."
        )
    return (
        "📋 The document is missing key details:\n\n"
        "1. Full name\n"
        "2. Email or another contact if you want it included\n\n"
        "Reply in one message."
    )

def generate_document_for_state(chat_id: int, state: dict, lang: str, doc_type: str):
    send_message(chat_id, TEXT[lang]["doc_thinking"])
    full_text = (state.get("case_text") or "")
    if state.get("doc_extra_data"):
        full_text += "\n\n--- PAPILDOMI DUOMENYS DOKUMENTUI / ADDITIONAL DOCUMENT DETAILS ---\n" + state.get("doc_extra_data", "")
    answer = openai_request(build_doc_prompt(full_text, lang, state.get("country") or "lt", doc_type), lang)
    if not answer:
        err = "⚠️ Dokumento šiuo metu nepavyko paruošti. Bandykite dar kartą po kelių minučių." if lang == "lt" else "⚠️ The document could not be prepared right now. Please try again in a few minutes."
        send_message(chat_id, err)
        return
    if answer.strip().upper().startswith("NEED_MORE_DATA"):
        details = answer.split(":", 1)[1].strip() if ":" in answer else ""
        msg = missing_doc_message(lang, details if details else "1. Vardas ir pavardė\n2. Data arba laikotarpis\n3. El. paštas arba telefonas")
        state["awaiting_doc_details"] = True
        state["pending_doc_type"] = doc_type
        send_message(chat_id, msg)
        return
    if document_has_placeholders(answer):
        state["awaiting_doc_details"] = True
        state["pending_doc_type"] = doc_type
        send_message(chat_id, placeholder_fix_message(lang))
        return
    answer = sanitize_generated_document(answer)
    send_chunks(chat_id, answer)


def case_list_title(lang: str) -> str:
    if lang == "lt":
        return "📁 Jūsų išsaugotos bylos"
    if lang == "no":
        return "📁 Dine lagrede saker"
    return "📁 Your saved cases"


def no_saved_cases_text(lang: str) -> str:
    if lang == "lt":
        return "📁 Išsaugotų bylų dar nėra. Galite pradėti naują bylą."
    if lang == "no":
        return "📁 Ingen lagrede saker ennå. Du kan starte en ny sak."
    return "📁 No saved cases yet. You can start a new case."


def case_loaded_text(lang: str, case_number: str | None) -> str:
    if lang == "lt":
        return f"✅ Byla atidaryta.\n\n🆔 Bylos Nr.: {case_number or ''}"
    if lang == "no":
        return f"✅ Saken er åpnet.\n\n🆔 Saksnr.: {case_number or ''}"
    return f"✅ Case opened.\n\n🆔 Case No.: {case_number or ''}"


def show_case_list(chat_id: int, lang: str):
    cases = db_list_cases(chat_id, limit=10)
    if not cases:
        send_message(chat_id, no_saved_cases_text(lang), reply_markup={"inline_keyboard": [[{"text": TEXT[lang]["new_case"], "callback_data": "new_case"}]]})
        return

    rows = []
    for c in cases:
        case_id = c.get("id")
        case_number = c.get("case_number") or f"CASE-{case_id}"
        title = (c.get("title") or "").replace("\n", " ").strip()
        if len(title) > 36:
            title = title[:33] + "..."
        button_text = f"{case_number} — {title}" if title else str(case_number)
        rows.append([{"text": button_text, "callback_data": f"select_case_{case_id}"}])
    rows.append([{"text": TEXT[lang]["new_case"], "callback_data": "new_case"}])
    send_message(chat_id, case_list_title(lang), reply_markup={"inline_keyboard": rows})


def load_case_into_state(chat_id: int, state: dict, lang: str, case_id: int):
    case = db_get_case(case_id, chat_id)
    if not case:
        send_message(chat_id, no_saved_cases_text(lang))
        return
    files = db_list_case_files(case_id)
    state.update({
        "case_text": case.get("case_text") or "",
        "case_id": case.get("id"),
        "case_number": case.get("case_number"),
        "awaiting_followup": False,
        "awaiting_doc_details": False,
        "pending_doc_type": None,
        "doc_extra_data": "",
        "files": files,
        "upload_menu_sent": False,
    })
    send_message(chat_id, case_loaded_text(lang, state.get("case_number")), reply_markup=file_action_menu(lang))


def build_followup_prompt(case_text: str, question: str, lang: str, country: str):
    country_name = COUNTRIES.get(country, COUNTRIES["lt"])[lang]
    if lang == "lt":
        return f"""
Atsakyk tik lietuviškai. Taikoma šalis / jurisdikcija: {country_name}.
Atsakyk į papildomą klausimą pagal bylą. Būk konkretus, aiškus ir praktiškas. Jei aktualu, nurodyk teisės akto straipsnį.

Bylos informacija:
{case_text}

Klausimas:
{question}
"""
    return f"""
Reply only in {lang_name(lang)}. Jurisdiction: {country_name}.
Answer the follow-up question based on the case. Be clear and practical. Mention legal sections if relevant.

Case information:
{case_text}

Question:
{question}
"""


def show_start(chat_id: int, tg_user: dict):
    lang = detect_lang(tg_user)
    state = get_state(chat_id)
    state.update({
        "lang": lang,
        "accepted": False,
        "country": None,
        "case_text": None,
        "case_id": None,
        "case_number": None,
        "awaiting_followup": False,
        "awaiting_doc_details": False,
        "pending_doc_type": None,
        "doc_extra_data": "",
        "files": [],
        "upload_menu_sent": False,
    })
    db_user_upsert(tg_user, lang)
    start_text = f"{welcome_text(lang)}\n\n{SAFETY_TEXT[lang]}"
    send_message(chat_id, start_text, reply_markup=safety_menu(lang))

def show_country(chat_id: int):
    lang = get_state(chat_id)["lang"]
    send_message(chat_id, TEXT[lang]["choose_country"], reply_markup=country_menu())


def show_collect_case(chat_id: int):
    state = get_state(chat_id)
    send_message(chat_id, TEXT[state["lang"]]["collect_case"])


def reset_current_case(state: dict):
    # Keep language, accepted safety status, selected country and subscription.
    # Clear only the active in-memory case selection. Existing cases stay saved in Supabase and can be opened from “My cases”.
    state.update({
        "case_text": None,
        "case_id": None,
        "case_number": None,
        "awaiting_followup": False,
        "awaiting_doc_details": False,
        "pending_doc_type": None,
        "doc_extra_data": "",
        "files": [],
        "upload_menu_sent": False,
    })


def start_new_case(chat_id: int, state: dict, lang: str):
    active, until = db_subscription_active(chat_id)
    if not active:
        reset_current_case(state)
        send_subscription_required(chat_id, lang)
        return
    reset_current_case(state)
    send_message(chat_id, TEXT[lang]["new_case_started"])
    show_collect_case(chat_id)


def create_case_from_text(chat_id: int, user_text: str):
    state = get_state(chat_id)
    lang = state["lang"]
    country = state.get("country") or "lt"
    case_id, case_number = db_create_case(chat_id, user_text)
    state.update({
        "case_text": user_text,
        "case_id": case_id,
        "case_number": case_number,
        "awaiting_followup": False,
        "doc_extra_data": "",
        "files": [],
        "upload_menu_sent": False,
    })
    send_message(chat_id, TEXT[lang]["case_created"].format(case_number=case_number))

    active, until = db_subscription_active(chat_id)
    if active:
        send_message(chat_id, f"{TEXT[lang]['active_until']} {until.date()}")
        send_message(chat_id, TEXT[lang]["thinking"])
        full_text = f"Bylos numeris: {case_number}\n\n{user_text}"
        answer = normalize_case_answer(openai_request(build_case_prompt(full_text, lang, country, full=True), lang))
        if not answer:
            send_message(chat_id, "⚠️ Analizės šiuo metu nepavyko atlikti. Bandykite dar kartą po kelių minučių." if lang == "lt" else "⚠️ The analysis could not be completed right now. Please try again in a few minutes.")
            return
        send_chunks(chat_id, answer, reply_markup=after_full_menu(lang))
        return

    send_message(chat_id, TEXT[lang]["thinking"])
    answer = normalize_case_answer(openai_request(build_case_prompt(user_text, lang, country, full=False), lang))
    if not answer:
        send_message(chat_id, "⚠️ Analizės šiuo metu nepavyko atlikti. Bandykite dar kartą po kelių minučių." if lang == "lt" else "⚠️ The analysis could not be completed right now. Please try again in a few minutes.")
        return
    send_chunks(chat_id, answer)
    send_message(chat_id, TEXT[lang]["initial"])
    send_message(chat_id, TEXT[lang]["choose_plan"], reply_markup=plan_menu(lang))


def handle_file(chat_id: int, msg: dict):
    state = get_state(chat_id)
    lang = state["lang"]
    media_group_id = msg.get("media_group_id")
    if media_group_id:
        MEDIA_GROUPS.setdefault(media_group_id, {"chat_id": chat_id, "count": 0})
        MEDIA_GROUPS[media_group_id]["count"] += 1

    file_id = None
    file_name = "uploaded_file"
    file_type = "unknown"
    caption = (msg.get("caption") or "").strip()

    if msg.get("document"):
        doc = msg["document"]
        file_id = doc.get("file_id")
        file_name = doc.get("file_name") or "document"
        file_type = doc.get("mime_type") or "document"
    elif msg.get("photo"):
        photo = msg["photo"][-1]
        file_id = photo.get("file_id")
        file_name = "photo.jpg"
        file_type = "image/jpeg"

    if not file_id:
        send_message(chat_id, TEXT[lang]["file_no_text"])
        return

    f_res = requests.get(f"{TELEGRAM_API}/getFile", params={"file_id": file_id}, timeout=20).json()
    file_path = f_res.get("result", {}).get("file_path")
    file_bytes = b""
    if file_path:
        file_bytes = requests.get(f"https://api.telegram.org/file/bot{TELEGRAM_BOT_TOKEN}/{file_path}", timeout=60).content

    extracted = extract_text_from_file(file_name, file_bytes)

    created_now = False
    if not state.get("case_id"):
        base = f"Byla sukurta iš dokumento: {file_name}" if lang == "lt" else f"Case created from document: {file_name}"
        if caption:
            base += f"\n\nVartotojo komentaras: {caption}" if lang == "lt" else f"\n\nUser comment: {caption}"
        case_id, case_number = db_create_case(chat_id, base)
        state.update({
            "case_id": case_id,
            "case_number": case_number,
            "case_text": base,
            "doc_extra_data": "",
            "files": [],
            "upload_menu_sent": False,
        })
        created_now = True

    db_add_case_file(state["case_id"], file_name, file_type, file_id)
    state.setdefault("files", []).append(file_name)

    additions = []
    if caption:
        additions.append(f"--- VARTOTOJO KOMENTARAS PRIE FAILO {file_name} ---\n{caption}")
    if extracted:
        additions.append(f"--- FILE: {file_name} ---\n{extracted[:6000]}")

    if additions:
        state["case_text"] = (state.get("case_text") or "") + "\n\n" + "\n\n".join(additions)
        db_update_case(state["case_id"], state["case_text"])

    if lang == "lt":
        if created_now:
            status_text = (
                f"📁 Byla sukurta\n\n"
                f"🆔 Bylos Nr.: {state.get('case_number')}\n\n"
                f"📎 Dokumentas pridėtas prie bylos."
            )
        else:
            status_text = "📎 Dokumentas pridėtas prie bylos."
        status_text += "\n\n✅ Galite pradėti bendrą bylos analizę."
    elif lang == "no":
        if created_now:
            status_text = (
                f"📁 Sak opprettet\n\n"
                f"🆔 Saksnr.: {state.get('case_number')}\n\n"
                f"📎 Dokumentet er lagt til saken."
            )
        else:
            status_text = "📎 Dokumentet er lagt til saken."
        status_text += "\n\n✅ Du kan starte samlet analyse."
    else:
        if created_now:
            status_text = (
                f"📁 Case created\n\n"
                f"🆔 Case No.: {state.get('case_number')}\n\n"
                f"📎 Document added to the case."
            )
        else:
            status_text = "📎 Document added to the case."
        status_text += "\n\n✅ You can start the combined case analysis."

    if media_group_id and state.get("upload_menu_sent"):
        # Telegram sends each file in an album as a separate webhook update.
        # After the first album/file message, keep later album files silent to avoid repeated messages.
        return

    if not state.get("upload_menu_sent"):
        send_message(chat_id, status_text, reply_markup=file_action_menu(lang))
        state["upload_menu_sent"] = True
    else:
        # Subsequent files are added silently to avoid repeated messages.
        pass


@app.route("/", methods=["GET"])
def health():
    return jsonify({"status": "ok", "service": "Justice AI PRO", "telegram": bool(TELEGRAM_BOT_TOKEN), "supabase": bool(SUPABASE_URL)})


@app.route("/telegram-webhook", methods=["POST"])
def telegram_webhook():
    update = request.get_json(force=True, silent=True) or {}
    try:
        if "callback_query" in update:
            cq = update["callback_query"]
            callback_id = cq.get("id")
            data = cq.get("data", "")
            msg = cq.get("message", {})
            chat_id = msg.get("chat", {}).get("id")
            user = cq.get("from", {})
            answer_callback(callback_id)
            if not chat_id:
                return jsonify({"ok": True})

            state = get_state(chat_id)
            lang = state.get("lang") or detect_lang(user)

            if data == "accept_safety":
                state["accepted"] = True
                show_country(chat_id)

            elif data.startswith("country_"):
                country = data.replace("country_", "")
                lang = language_for_country_selection(country, state.get("lang"), user)
                state["lang"] = lang
                state["country"] = country
                db_user_upsert(user, lang)

                active, until = db_subscription_active(chat_id)
                if active:
                    send_message(chat_id, f"{TEXT[lang]['active_until']} {until.date()}")
                    if state.get("case_id") and state.get("case_text"):
                        send_message(chat_id, "✅ Galite tęsti aktyvią bylą." if lang == "lt" else ("✅ Du kan fortsette den aktive saken." if lang == "no" else "✅ You can continue the active case."), reply_markup=file_action_menu(lang))
                    else:
                        show_case_list(chat_id, lang)
                        show_collect_case(chat_id)
                else:
                    send_subscription_required(chat_id, lang)
                return jsonify({"ok": True})

            elif data.startswith("plan_"):
                plan_key = data.replace("plan_", "")
                plan = PLANS.get(plan_key)
                if not plan:
                    return jsonify({"ok": True})
                description = f"{plan[lang]} access"
                payload = f"sub|{plan_key}|{chat_id}"
                requests.post(f"{TELEGRAM_API}/sendInvoice", json={
                    "chat_id": chat_id,
                    "title": "Justice AI Access",
                    "description": description,
                    "payload": payload,
                    "provider_token": "",
                    "currency": "XTR",
                    "prices": [{"label": description, "amount": plan["stars"]}],
                }, timeout=20)

            elif data == "ask_question":
                if not state.get("case_text"):
                    send_message(chat_id, TEXT[lang]["no_case"])
                else:
                    state["awaiting_followup"] = True
                    send_message(chat_id, TEXT[lang]["ask_prompt"])

            elif data == "new_case":
                start_new_case(chat_id, state, lang)

            elif data == "case_list":
                show_case_list(chat_id, lang)

            elif data.startswith("select_case_"):
                try:
                    selected_case_id = int(data.replace("select_case_", ""))
                    load_case_into_state(chat_id, state, lang, selected_case_id)
                except Exception:
                    show_case_list(chat_id, lang)

            elif data == "upload_more":
                # Legacy callback kept for old Telegram messages. New menus no longer show this button.
                if lang == "lt":
                    send_message(chat_id, "📎 Dokumentą galite pridėti per Telegram prisegimo ikoną.", reply_markup=file_action_menu(lang))
                elif lang == "no":
                    send_message(chat_id, "📎 Du kan legge til dokumenter med vedlegg-ikonet i Telegram.", reply_markup=file_action_menu(lang))
                else:
                    send_message(chat_id, "📎 You can add documents using the Telegram attachment icon.", reply_markup=file_action_menu(lang))

            elif data == "case_docs":
                files = state.get("files") or []
                if not files:
                    send_message(chat_id, "📂 Bylos dokumentų dar nėra." if lang == "lt" else "📂 No case documents yet.", reply_markup=file_action_menu(lang))
                else:
                    lines = [f"{i + 1}. {name}" for i, name in enumerate(files)]
                    title = "📂 Bylos dokumentai:" if lang == "lt" else ("📂 Saksdokumenter:" if lang == "no" else "📂 Case documents:")
                    send_message(chat_id, title + "\n\n" + "\n".join(lines), reply_markup=file_action_menu(lang))

            elif data == "analyze_case":
                if not state.get("case_text"):
                    send_message(chat_id, TEXT[lang]["no_case"])
                else:
                    active, until = db_subscription_active(chat_id)
                    if active:
                        send_message(chat_id, TEXT[lang]["thinking"])
                        full_text = f"Bylos numeris: {state.get('case_number')}\n\n{state.get('case_text')}"
                        answer = normalize_case_answer(openai_request(build_case_prompt(full_text, lang, state.get("country") or "lt", full=True), lang))
                        if not answer:
                            err = "⚠️ Analizės šiuo metu nepavyko atlikti. Bandykite dar kartą po kelių minučių." if lang == "lt" else "⚠️ The analysis could not be completed right now. Please try again in a few minutes."
                            send_message(chat_id, err)
                        else:
                            send_chunks(chat_id, answer, reply_markup=after_full_menu(lang))
                    else:
                        send_message(chat_id, TEXT[lang]["initial"])
                        send_message(chat_id, TEXT[lang]["choose_plan"], reply_markup=plan_menu(lang))

            elif data == "show_docs":
                if not state.get("case_text"):
                    send_message(chat_id, TEXT[lang]["no_case"])
                else:
                    send_message(chat_id, TEXT[lang]["choose_doc"], reply_markup=doc_menu(lang, state.get("case_text") or "", state.get("country")))

            elif data.startswith("doc_"):
                if not state.get("case_text"):
                    send_message(chat_id, TEXT[lang]["no_case"])
                else:
                    doc_type = data.replace("doc_", "")
                    full_text = state.get('case_text') or ""
                    check = openai_request(build_doc_missing_prompt(full_text, lang, state.get("country") or "lt", doc_type), lang)
                    missing = parse_missing_doc_questions(check)
                    if missing:
                        state["awaiting_doc_details"] = True
                        state["pending_doc_type"] = doc_type
                        send_message(chat_id, missing_doc_message(lang, missing))
                    else:
                        generate_document_for_state(chat_id, state, lang, doc_type)

            return jsonify({"ok": True})

        if "pre_checkout_query" in update:
            pcq = update["pre_checkout_query"]
            requests.post(f"{TELEGRAM_API}/answerPreCheckoutQuery", json={"pre_checkout_query_id": pcq.get("id"), "ok": True}, timeout=10)
            return jsonify({"ok": True})

        if "message" in update:
            msg = update["message"]
            chat_id = msg.get("chat", {}).get("id")
            user = msg.get("from", {})
            text = msg.get("text", "")
            if not chat_id:
                return jsonify({"ok": True})
            state = get_state(chat_id)

            if text.startswith("/start"):
                show_start(chat_id, user)
                return jsonify({"ok": True})

            if text.startswith("/newcase") or text.startswith("/naujabyla"):
                lang = state.get("lang") or detect_lang(user)
                if not state.get("accepted"):
                    show_start(chat_id, user)
                elif not state.get("country"):
                    show_country(chat_id)
                else:
                    start_new_case(chat_id, state, lang)
                return jsonify({"ok": True})

            if text.startswith("/cases") or text.startswith("/bylos"):
                lang = state.get("lang") or detect_lang(user)
                if not state.get("accepted"):
                    show_start(chat_id, user)
                elif not state.get("country"):
                    show_country(chat_id)
                else:
                    show_case_list(chat_id, lang)
                return jsonify({"ok": True})

            if msg.get("successful_payment"):
                payload = msg["successful_payment"].get("invoice_payload", "")
                if payload.startswith("sub|"):
                    _, plan_key, _ = payload.split("|")
                    plan = PLANS.get(plan_key)
                    if plan:
                        lang = state.get("lang") or detect_lang(user)
                        until = db_set_subscription(chat_id, plan_key, plan["months"], plan["stars"])
                        send_message(chat_id, f"{TEXT[lang]['access_active']}\n{TEXT[lang]['active_until']} {until.date()}")
                        reset_current_case(state)
                        show_collect_case(chat_id)
                return jsonify({"ok": True})

            if msg.get("document") or msg.get("photo"):
                if not state.get("accepted"):
                    show_start(chat_id, user)
                elif not state.get("country"):
                    show_country(chat_id)
                else:
                    active, until = db_subscription_active(chat_id)
                    if not active:
                        send_subscription_required(chat_id, lang)
                    else:
                        handle_file(chat_id, msg)
                return jsonify({"ok": True})

            lang = state.get("lang") or detect_lang(user)
            if not state.get("accepted"):
                show_start(chat_id, user)
            elif not state.get("country"):
                show_country(chat_id)
            elif not db_subscription_active(chat_id)[0]:
                send_subscription_required(chat_id, lang)
            elif state.get("awaiting_doc_details") and state.get("case_text"):
                state["awaiting_doc_details"] = False
                doc_type = state.get("pending_doc_type") or "authority_complaint"
                state["pending_doc_type"] = None
                state["case_text"] = (state.get("case_text") or "") + f"\n\n--- PAPILDOMI DUOMENYS DOKUMENTUI ---\n{text}"
                if state.get("case_id"):
                    db_update_case(state["case_id"], state["case_text"])
                generate_document_for_state(chat_id, state, lang, doc_type)
            elif state.get("awaiting_followup") and state.get("case_text"):
                state["awaiting_followup"] = False
                send_message(chat_id, TEXT[lang]["thinking"])
                answer = openai_request(build_followup_prompt(state["case_text"], text, lang, state.get("country") or "lt"), lang)
                send_chunks(chat_id, answer, reply_markup=after_full_menu(lang))
            else:
                if state.get("case_id") and state.get("case_text"):
                    state["case_text"] = (state.get("case_text") or "") + f"\n\n--- PAPILDOMA VARTOTOJO INFORMACIJA ---\n{text}"
                    db_update_case(state["case_id"], state["case_text"])
                    if lang == "lt":
                        send_message(chat_id, "✅ Informacija pridėta prie bylos.\n\nGalite pradėti bendrą bylos analizę.", reply_markup=file_action_menu(lang))
                    elif lang == "no":
                        send_message(chat_id, "✅ Informasjonen er lagt til saken.\n\nDu kan starte samlet analyse.", reply_markup=file_action_menu(lang))
                    else:
                        send_message(chat_id, "✅ Information added to the case.\n\nYou can start the combined case analysis.", reply_markup=file_action_menu(lang))
                else:
                    create_case_from_text(chat_id, text)

            return jsonify({"ok": True})

    except Exception as e:
        logger.exception("Webhook error: %s", e)
        return jsonify({"ok": False, "error": str(e)}), 500

    return jsonify({"ok": True})


@app.route("/set-webhook", methods=["GET"])
def set_webhook():
    if not PUBLIC_BASE_URL:
        return jsonify({"ok": False, "error": "PUBLIC_BASE_URL or RENDER_EXTERNAL_URL missing"}), 400
    webhook_url = PUBLIC_BASE_URL.rstrip("/") + "/telegram-webhook"
    r = requests.post(f"{TELEGRAM_API}/setWebhook", json={"url": webhook_url}, timeout=20)
    return jsonify(r.json())


if __name__ == "__main__":
    port = int(os.getenv("PORT", "5000"))
    app.run(host="0.0.0.0", port=port)
